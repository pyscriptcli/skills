#!/usr/bin/env python3
"""
Diagnostic utility to inspect slides, shapes, tables, and images of any PPTX file.
"""

import sys
import pptx
from pathlib import Path

def inspect(file_path):
    prs = pptx.Presentation(file_path)
    print(f"Presentation: {file_path}")
    print(f"Total Slides: {len(prs.slides)}")
    print(f"Dimensions: {prs.slide_width/914400:.2f} x {prs.slide_height/914400:.2f} inches")

    for i, slide in enumerate(prs.slides):
        print(f"\n--- Slide {i+1} (ID: {slide.slide_id}) ---")
        for idx, shp in enumerate(slide.shapes):
            dim = f"L:{shp.left/914400:.2f}\" T:{shp.top/914400:.2f}\" W:{shp.width/914400:.2f}\" H:{shp.height/914400:.2f}\""
            txt = ""
            if shp.has_text_frame:
                txt = " | ".join([p.text.strip() for p in shp.text_frame.paragraphs if p.text.strip()])
            elif shp.has_table:
                tbl = shp.table
                txt = f"Table ({len(tbl.rows)}x{len(tbl.columns)}): " + " / ".join([c.text.strip() for c in tbl.rows[0].cells])
            elif shp.shape_type == pptx.enum.shapes.MSO_SHAPE_TYPE.PICTURE:
                txt = f"Picture ({shp.image.content_type})"
            print(f"  [{idx}] {shp.name} ({shp.shape_type}) [{dim}] => {txt[:80]}")

if __name__ == "__main__":
    path = sys.argv[1] if len(sys.argv) > 1 else "templates/prime_proposal_template.pptx"
    inspect(path)
